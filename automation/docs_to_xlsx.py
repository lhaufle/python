import docx #import module to read or write to word doc
from openpyxl import Workbook #import module and class to create an excel document

#get the name of the file
doc_name = input('Please enter the name of the document:')

try:
    #open work document to read
    doc = docx.Document("add your directory/" + doc_name)
except docx.opc.exceptions.PackageNotFoundError:
    print("The File could not be found, please check directory and run again!")

'''
if the location of the data you found in your document is static, 
then read from the specific paragraph.
'''
all_p = doc.paragraphs[5]

#read the int values from doc and turn them into a list
values = list(all_p.text.split(' '))

#check values to make sure it works
print("Exporting values:")
print(values)

#create an excel workbook
workbook = Workbook()

#open worksheet to edit
sheet = workbook.active

#create a header
sheet["A1"] = "Units Sold"

#add values below the header
count = 2 #count starts at 2 becuase values must be below the header in record A1
for value in values:
    sheet["A" + str(count)] = value
    count += 1

#add the totals
total_value = 0
for val in values:
    total_value += int(val)

#write the total into excel below the listed values
sheet["A" + str(count)] = total_value

#save the worksheet with values
workbook.save(filename="add directory and file name")

#confirm job complete
print("Export Complete")
